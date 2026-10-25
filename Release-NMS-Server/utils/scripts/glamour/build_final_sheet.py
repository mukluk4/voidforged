# Writes glamour_final.xlsx: every "Glamour - '<item>'" in the database, Y if a vendor sells it (and which),
# N with the reason it was left out. Uses gen_glamour_vendors.py's own rules, so it matches the vendor SQL.
#   python build_final_sheet.py <eqemu_config.json> <glamour_review.xlsx> <out.xlsx>
import os, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

config, review, out = sys.argv[1:4]
here = os.path.dirname(os.path.abspath(__file__))

# Run the generator's list building (everything before its output modes) with the review sheet applied
src = open(os.path.join(here, 'gen_glamour_vendors.py')).read().split("if mode == 'json':")[0]
g = {'__name__': 'gen_glamour_vendors'}
sys.argv = ['gen_glamour_vendors.py', config, 'none', '--review', review]
exec(compile(src, 'gen_glamour_vendors.py', 'exec'), g)
query, SLOTS, ARMOR_TYPES, ROBES = g['query'], g['SLOTS'], g['ARMOR_TYPES'], g['ROBES']
sold = {gid: vendor for vendor, ids in g['vendors'] for gid in ids}
armor_pick = {gid: key for key, (_r, gid) in g['best'].items()}
weapon_pick = {key: gid for key, (_r, gid) in g['best_models'].items()}
armor_rep = {key: gid for key, (_r, gid) in g['best'].items()}

# Review sheet: why a row was set to N (its Flag, or that you set it), and the Patterns rules
wb_review = load_workbook(review, read_only=True)
review_flag = {}
for ws in wb_review.worksheets:
    if ws.title in ('How to use', 'Summary', 'Patterns'):
        continue
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None and str(row[4]).strip().upper() == 'N':
            review_flag[int(row[0])] = f'review: {row[3]}' if row[3] else 'review: set to N'
rules = [(str(p), (str(v).strip() if v else '')) for p, v, *_ in
         wb_review['Patterns'].iter_rows(min_row=2, values_only=True) if p]

src_type = {name: int(t) for name, t in query(
    "SELECT Name, MIN(itemtype) FROM items WHERE itemtype IN (0,1,2,3,4,5,7,8,35,45) AND Name IN ("
    "SELECT SUBSTRING(Name, 12, CHAR_LENGTH(Name) - 12) FROM items WHERE Name LIKE 'Glamour - %' "
    "AND (slots & 26624)) GROUP BY Name")}

rows = []
for gid, name, slots, material, color, hf, idfile in query(
        "SELECT id, Name, slots, material, color, herosforgemodel, idfile FROM items "
        "WHERE Name LIKE 'Glamour - %' ORDER BY id"):
    gid, slots, material, color, hf = int(gid), int(slots), int(material), int(color), int(hf)
    look = f'texture {material}, tint {color & 0xFFFFFF:06X}'
    if slots & 26624:
        vendor = g['WEAPON_TYPES'].get(src_type.get(name[11:-1]), 'Unusual Weapons')
        look = f'model {idfile}'
        if idfile == 'IT0':
            why = 'no model (IT0)'
        elif idfile == 'IT63':
            why = 'dropped-item bag model (IT63)'
        elif weapon_pick.get((vendor, idfile)) != gid:
            why = f'same model as {weapon_pick[(vendor, idfile)]}'
        else:
            why = ''
    else:
        slot_name = next((n for bit, n in SLOTS if slots & bit), None)
        if not slot_name:
            vendor, why = '', 'not a vendor slot (face, ear, neck, ...)'
        else:
            kind = 'Robes' if slot_name == 'Chests' and material in ROBES else ARMOR_TYPES.get(material, 'Exotic')
            vendor = kind if kind == 'Robes' or kind in g['MERGED_TYPES'] else f'{kind} {slot_name}'
            key = (kind, slot_name, material, color)
            if hf:
                why, look = "Hero's Forge model", f"Hero's Forge {hf}"
            elif material in g['SKIP_MATERIALS']:
                why = f'texture {material} not verified on the RoF2 client'
            elif armor_rep.get(key) != gid:
                why = f'same look as {armor_rep[key]}'
            elif color not in g['tint_order'][material]:
                why = 'extra tint (not one of the 8 matching colors)'
            else:
                why = ''
    if not why and gid not in sold:
        why = review_flag.get(gid) or next(
            (f'pattern: "{p}"' for p, v in rules if p.lower() in name.lower() and (not v or v == vendor)), 'review')
    rows.append((gid, name, vendor if not why else (vendor or ''), look,
                 'Y' if not why else 'N', sold.get(gid, '') if not why else why))

FONT, BOLD = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True)
HEAD = PatternFill('solid', start_color='D9D9D9')
YES, NO = PatternFill('solid', start_color='D9EAD3'), PatternFill('solid', start_color='F4CCCC')

wb = Workbook()
summary = wb.active
summary.title = 'Summary'
ws = wb.create_sheet('All glamours')
headers = ['Item ID', 'Glamour', 'Vendor group', 'Look', 'On vendor', 'Vendor / reason']
for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font, c.fill = BOLD, HEAD
order = {v: i for i, (v, _) in enumerate(g['vendors'])}
rows.sort(key=lambda r: (r[4] != 'Y', order.get(r[2], len(order)), r[2], r[0]))
for r, row in enumerate(rows, 2):
    for col, val in enumerate(row, 1):
        ws.cell(row=r, column=col, value=val).font = FONT
    ws.cell(row=r, column=5).fill = YES if row[4] == 'Y' else NO
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:F{len(rows) + 1}'
for col, w in zip('ABCDEF', (10, 56, 18, 26, 10, 52)):
    ws.column_dimensions[col].width = w

# Summary: counts by reason (formulas, so a filter or edit on the main tab can be checked against it)
reasons = {}
for row in rows:
    key = 'On a vendor' if row[4] == 'Y' else row[5].split(' as ')[0].split(':')[0]
    reasons[key] = reasons.get(key, 0) + 1
summary.append(['Glamour list: every glamour in the database, on a vendor (Y) or left out (N, with the reason)'])
summary['A1'].font = BOLD
summary.append([])
summary.append(['Outcome', 'Items'])
for c in summary[3]:
    c.font, c.fill = BOLD, HEAD
for key, n in sorted(reasons.items(), key=lambda kv: -kv[1]):
    summary.append([key, n])
last = summary.max_row
summary.append(['Total', f'=SUM(B4:B{last})'])
summary.append(['Check: rows on the list', f"=COUNTA('All glamours'!A:A)-1"])
summary.append(['Check: Y on the list', f"=COUNTIF('All glamours'!E:E,\"Y\")"])
for row in summary.iter_rows(min_row=4):
    for c in row:
        c.font = BOLD if c.row == last + 1 else FONT
summary.column_dimensions['A'].width = 52
summary.column_dimensions['B'].width = 10

wb.save(out)
print('saved', out, len(rows), 'glamours,', sum(r[4] == 'Y' for r in rows), 'on vendors')
for key, n in sorted(reasons.items(), key=lambda kv: -kv[1]):
    print(f'  {n:6}  {key}')
